"""Write Claude's title/abstract decisions for the 50 pilot records into 03_screening.xlsx.

Usage: python3 -I pilot_decisions_claude.py ../03_screening.xlsx
Basis: title + abstract only (screening stage). E9 is a proposed new code (see protocol v0.3 proposal).
"""
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

D = {
    "R00745": ("EXCLUDE", "E2", "Tear-flow (Sno) strips used to collect samples for HPLC drug assay; not an LFA, no virus target."),
    "R03476": ("INCLUDE", "", "EQA data on commercial EBV (infectious mononucleosis) POC tests, mostly immunochromatographic; commercial-evaluation group. Note: heterophile antibody target."),
    "R00616": ("EXCLUDE", "E6", "Physician survey on influenza management; no assay data."),
    "R08180": ("INCLUDE", "", "RAA + lateral flow dipstick, triplex bovine viruses (BCoV, IBRV, BVDV), field swabs. Label not stated: check full text."),
    "R02326": ("EXCLUDE", "E9", "Outbreak investigation; RDT mentioned as trigger, testing by ELISA/PCR; no LFA evaluated."),
    "R04307": ("EXCLUDE", "E6", "Public survey on LFT vs PCR use; no assay data."),
    "R00119": ("EXCLUDE", "E2", "RIBA strip immunoblot for HCV serotyping; immunoblot, not lateral flow."),
    "R07527": ("EXCLUDE", "E4", "Narrative review of TB diagnostics (also not a virus target)."),
    "R03821": ("INCLUDE", "", "Accuracy of two SARS-CoV-2 IgG/IgM LFIAs vs CLIA, 259 samples; commercial-evaluation group."),
    "R03547": ("INCLUDE", "", "Two SARS-CoV-2 antigen tests; one is fluorescent (FIA, out of scope), one visual RAT (in scope); commercial-evaluation group."),
    "R00545": ("UNSURE", "", "Directigen Flu A+B 'enzyme immunoassay' (rapid membrane test) vs DFA; unclear if lateral flow; check full text."),
    "R00022": ("EXCLUDE", "E5", "Portuguese; review on adenosine deaminase (also E4)."),
    "R03488": ("UNSURE", "", "Letter, no abstract; reports clinical use of a SARS-CoV-2 rapid antigen test. May contain original data; full text needed."),
    "R07524": ("EXCLUDE", "E4", "Narrative review (pharmacist POCT)."),
    "R06673": ("EXCLUDE", "E2", "Electrochemical paper-based biosensor for mpox; not a lateral flow assay."),
    "R05781": ("EXCLUDE", "E9", "Healthcare-worker infection epidemiology; RATs used as a tool, not evaluated."),
    "R03344": ("INCLUDE", "", "Dual-mode LFIA for SARS-CoV-2 IgA: colorimetric (nanogold, smartphone) + chemiluminescence; colorimetric channel in scope."),
    "R03517": ("INCLUDE", "", "RT-RPA for SARS-CoV-2 with lateral flow strip readout, 113 clinical samples."),
    "R07599": ("UNSURE", "", "Feasibility/acceptability of Ag-RDT testing in markets; no accuracy data. RULE NEEDED for implementation studies."),
    "R01414": ("EXCLUDE", "E1", "Colloidal gold strip for Mycoplasma suis (bacterium)."),
    "R03441": ("EXCLUDE", "E3", "Lateral flow with fluorescent-nanoparticle label only."),
    "R00620": ("EXCLUDE", "E5", "Turkish; Leishmania serology (also E1)."),
    "R01368": ("EXCLUDE", "E4", "Comment/letter on another paper; no abstract."),
    "R03214": ("EXCLUDE", "E1", "Cryptococcal antigen LFA (fungus)."),
    "R07991": ("INCLUDE", "", "SARS-CoV-2 in wastewater: visually read carbon black dipsticks (in scope) + fluorescent nanodiamond (out of scope)."),
    "R00928": ("EXCLUDE", "E5", "German case report; urine test strip, IFA/ELISA (also E2)."),
    "R08802": ("UNSURE", "", "Epitope mapping of antibodies from 43 commercial SARS-CoV-2 RATs; no strip testing. RULE NEEDED for component-only studies."),
    "R05681": ("EXCLUDE", "E4", "Review."),
    "R02553": ("EXCLUDE", "E5", "Chinese full text (English abstract): carbon nanoparticle LFA for Ebola. Relevant topic; excluded only by language rule."),
    "R02551": ("EXCLUDE", "E2", "Multiplex PCR (FilmArray) for meningitis; no LFA."),
    "R02319": ("EXCLUDE", "E2", "Multiplex PCR (FilmArray) panel; no LFA."),
    "R06459": ("EXCLUDE", "E3", "SERS-LFIA (Raman readout) only."),
    "R02883": ("EXCLUDE", "E9", "HIV prevalence study using oral rapid tests as a tool; test not evaluated."),
    "R04046": ("INCLUDE", "", "Evaluation of 20 serological SARS-CoV-2 LFAs; commercial-evaluation group."),
    "R08078": ("EXCLUDE", "E9", "Vaccination cohort; RAT used only to define infection."),
    "R02815": ("EXCLUDE", "E9", "STI prevalence using rapid immunochromatographic tests as a tool; not evaluated."),
    "R03878": ("EXCLUDE", "E1", "Urine LAM lateral flow for tuberculosis (bacterium)."),
    "R04777": ("EXCLUDE", "E1", "Malaria RDT in HIV-infected children (parasite target)."),
    "R06320": ("EXCLUDE", "E9", "Saliva RAT used to measure viral shedding vs blood counts; test not evaluated."),
    "R06842": ("EXCLUDE", "E5", "French; pulmonary embolism prognosis (also E9)."),
    "R05022": ("EXCLUDE", "E6", "Cost-effectiveness modelling of HCV programmes."),
    "R06029": ("EXCLUDE", "E9", "Sewage surveillance programme; RAT kits distributed, not evaluated."),
    "R02665": ("INCLUDE", "", "Colloidal gold immunochromatographic strip for Newcastle disease virus; Se/Sp vs RT-PCR."),
    "R03808": ("INCLUDE", "", "Colloidal gold immunochromatographic strip for H7 avian influenza; LOD in chicken swabs."),
    "R05444": ("EXCLUDE", "E1", "Survey on urine LAM (TB) test uptake (also E6)."),
    "R01636": ("UNSURE", "", "BD Veritor (reader-based 'digital immunoassay') vs Sofia (fluorescent, out of scope). Veritor label/readout to be checked."),
    "R05501": ("INCLUDE", "", "CRISPR-Cas13a for norovirus GII.4 with fluorescence and lateral flow strip readout; strip arm in scope."),
    "R07959": ("EXCLUDE", "E9", "Canine distemper treatment study; RDT used only for diagnosis."),
    "R06799": ("EXCLUDE", "E9", "Trial protocol; dengue RDT used as outcome measure."),
    "R01904": ("INCLUDE", "", "MERS-CoV nucleocapsid mAbs used to develop ELISA and a rapid immunochromatographic test; label to check in full text."),
}


def main(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["screening_all"]
    head = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(head)}
    done = 0
    for row in ws.iter_rows(min_row=2):
        rid = row[ci["rec_id"]].value
        if rid in D:
            dec, code, note = D[rid]
            row[ci["decision_claude"]].value = dec
            row[ci["reason_claude"]].value = code
            row[ci["note_claude"]].value = note
            done += 1
    assert done == len(D) == 50, done

    pilot = wb["pilot_author"]
    ph = [c.value for c in pilot[1]]
    pi = {h: i for i, h in enumerate(ph)}
    if "pilot_check" in wb.sheetnames:
        del wb["pilot_check"]
    chk = wb.create_sheet("pilot_check", 0)
    cols = ["rec_id", "pmid", "year", "title", "decision_claude", "reason_claude", "note_claude",
            "author_agrees (Y/N)", "author_comment"]
    chk.append(cols)
    for c in chk[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for row in pilot.iter_rows(min_row=2, values_only=True):
        rid = row[pi["rec_id"]]
        dec, code, note = D[rid]
        chk.append([rid, row[pi["pmid"]], row[pi["year"]], row[pi["title"]], dec, code, note, "", ""])
    for letter, w in zip("ABCDEFGHI", (9, 10, 6, 55, 11, 8, 70, 12, 30)):
        chk.column_dimensions[letter].width = w
    for r in chk.iter_rows(min_row=2):
        for c in (r[3], r[6]):
            c.alignment = Alignment(wrap_text=True, vertical="top")
    chk.freeze_panes = "B2"
    codes = wb["codes"]
    if not any(r[0].value == "E9" for r in codes.iter_rows(min_row=2)):
        codes.append(["E9", "PROPOSED: a virus LFA/RDT is only used as a tool (prevalence, cohort, treatment, "
                            "surveillance, trial outcome); no assay performance is evaluated"])
    wb.save(path)
    n = {k: sum(1 for v in D.values() if v[0] == k) for k in ("INCLUDE", "UNSURE", "EXCLUDE")}
    print(n)


if __name__ == "__main__":
    main(sys.argv[1])
