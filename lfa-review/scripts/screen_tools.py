"""Screening helpers.

  python3 -I screen_tools.py auto   WORKBOOK            apply deterministic rules (writes decisions/auto.tsv)
  python3 -I screen_tools.py batch  WORKBOOK N SIZE OUT  dump batch N of records still undecided to OUT
  python3 -I screen_tools.py merge  WORKBOOK            write all decisions/*.tsv into the workbook
  python3 -I screen_tools.py status WORKBOOK            counts

Decision files: decisions/*.tsv with lines  rec_id<TAB>DECISION<TAB>code<TAB>note
The master table (02_records/work/master_records.xlsx) supplies abstracts.
"""
import collections
import glob
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DEC_DIR = os.path.join(ROOT, "decisions")
MASTER = os.path.join(ROOT, "02_records", "work", "master_records.xlsx")

AUTO_E4 = {"Published Erratum", "Preprint", "Review", "Systematic Review", "Meta-Analysis", "Scoping Review",
           "Editorial", "News", "Newspaper Article", "Practice Guideline", "Guideline", "Historical Article",
           "Biography", "Portrait", "Interview"}


def load_master():
    ws = openpyxl.load_workbook(MASTER, read_only=True)["master"]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it]


def load_decisions():
    dec = {}
    for path in sorted(glob.glob(os.path.join(DEC_DIR, "*.tsv"))):
        for line in open(path, encoding="utf-8"):
            if not line.strip() or line.startswith("#"):
                continue
            parts = line.rstrip("\n").split("\t") + ["", "", ""]
            rid, d, code, note = parts[0].strip(), parts[1].strip().upper(), parts[2].strip(), parts[3].strip()
            assert d in {"INCLUDE", "EXCLUDE", "UNSURE"}, (path, line)
            dec[rid] = (d, code, note, os.path.basename(path))
    return dec


def auto(_wb):
    os.makedirs(DEC_DIR, exist_ok=True)
    rows = load_master()
    out = []
    for r in rows:
        pts = set((r["pub_types"] or "").split(";"))
        langs = (r["language"] or "").split(";")
        if "eng" not in langs:
            out.append((r["rec_id"], "EXCLUDE", "E5", "auto: language " + (r["language"] or "")))
        elif "Retracted Publication" in pts:
            out.append((r["rec_id"], "EXCLUDE", "E8", "auto: retracted publication"))
        elif pts & AUTO_E4 and "Letter" not in pts:
            out.append((r["rec_id"], "EXCLUDE", "E4", "auto: " + ";".join(sorted(pts & AUTO_E4))))
    with open(os.path.join(DEC_DIR, "auto.tsv"), "w", encoding="utf-8") as f:
        f.write("# deterministic rules, protocol decision log 2026-10-28\n")
        for o in out:
            f.write("\t".join(o) + "\n")
    print("auto decisions:", len(out), collections.Counter(o[2] for o in out))


def batch(_wb, n, size, out_path):
    n, size = int(n), int(size)
    dec = load_decisions()
    todo = [r for r in load_master() if r["rec_id"] not in dec]
    chunk = todo[:size]  # always the next undecided records; n only labels the batch
    with open(out_path, "w", encoding="utf-8") as f:
        for r in chunk:
            ab = (r["abstract"] or "(no abstract)").replace("\n", " ")
            if len(ab) > 900:
                ab = ab[:900] + " …"
            f.write(f"{r['rec_id']} | {r['year']} | {r['pub_types']}\nT: {r['title']}\nA: {ab}\n\n")
    print(f"remaining undecided: {len(todo)}; wrote {len(chunk)} to {out_path}; "
          f"range {chunk[0]['rec_id'] if chunk else '-'}..{chunk[-1]['rec_id'] if chunk else '-'}")


def merge(wb_path):
    dec = load_decisions()
    wb = openpyxl.load_workbook(wb_path)
    ws = wb["screening_all"]
    head = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(head)}
    for row in ws.iter_rows(min_row=2):
        rid = row[ci["rec_id"]].value
        if rid in dec and not row[ci["decision_claude"]].value:
            d, code, note, _ = dec[rid]
            row[ci["decision_claude"]].value = d
            row[ci["reason_claude"]].value = code
            row[ci["note_claude"]].value = note
    wb.save(wb_path)
    status(wb_path)


def status(wb_path):
    ws = openpyxl.load_workbook(wb_path, read_only=True)["screening_all"]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    ci = {h: i for i, h in enumerate(head)}
    c = collections.Counter()
    codes = collections.Counter()
    total = 0
    for r in it:
        total += 1
        d = r[ci["decision_claude"]]
        c[d or "undecided"] += 1
        if d == "EXCLUDE":
            codes[r[ci["reason_claude"]]] += 1
    print("total", total, dict(c), "exclusion codes", dict(sorted(codes.items())))


if __name__ == "__main__":
    cmd, args = sys.argv[1], sys.argv[2:]
    {"auto": auto, "batch": batch, "merge": merge, "status": status}[cmd](*args)
